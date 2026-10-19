"""
Excel, Word and Python File Processing
--------------------------------------
Run:
    pip install -r requirements.txt
    python process_files.py

The script reads registration_data.xlsx, cleans and organizes the data,
saves cleaned_registration_data.xlsx, and creates registration_summary.docx.
It does not require opening Excel or Word manually.
"""
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from docx.shared import Pt, RGBColor, Inches

BASE_DIR = Path(__file__).resolve().parent
INPUT_EXCEL = BASE_DIR / "registration_data.xlsx"
OUTPUT_EXCEL = BASE_DIR / "cleaned_registration_data.xlsx"
OUTPUT_WORD = BASE_DIR / "registration_summary.docx"

def clean_registration_data():
    df = pd.read_excel(INPUT_EXCEL, sheet_name="Registrations")
    original_count = len(df)

    # Normalize column names and text values
    df.columns = [str(col).strip().replace(" ", "_") for col in df.columns]
    for col in df.select_dtypes(include="object").columns:
        df[col] = df[col].astype("string").str.strip()

    # Normalize email addresses and department labels
    if "Email" in df.columns:
        df["Email"] = df["Email"].str.lower()
        df["Email"] = df["Email"].replace("", pd.NA)
    if "Department" in df.columns:
        df["Department"] = df["Department"].str.upper()
        df["Department"] = df["Department"].replace({
            "INFORMATION TECHNOLOGY": "IT",
            "CSE": "CSE",
            "ECE": "ECE",
            "AI & DS": "AI & DS",
            "IT": "IT",
        })

    # Remove duplicate registrations based on email, keeping first occurrence.
    before_dedup = len(df)
    valid_emails = df["Email"].notna()
    duplicates = df[valid_emails].duplicated(subset=["Email"], keep="first")
    duplicate_count = int(duplicates.sum())
    df = df.loc[~(valid_emails & duplicates)].copy()

    # Sort records by department, then name
    df = df.sort_values(["Department", "Full_Name"], na_position="last").reset_index(drop=True)

    # Add a review flag for missing emails
    df["Review_Status"] = df["Email"].apply(
        lambda value: "Email missing - review" if pd.isna(value) or str(value).strip() == "" else "OK"
    )

    with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Cleaned Registrations")
        ws = writer.sheets["Cleaned Registrations"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5597")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = min(max_length + 3, 34)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center")

    return df, original_count, duplicate_count

def create_word_report(df, original_count, duplicate_count):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    title = doc.add_heading("Student Registration Data Cleaning Report", 0)
    title.runs[0].font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph(
        "This report was generated automatically using Python. "
        "The source Excel workbook was read, cleaned, organized, and saved without manually opening Excel or Word."
    )

    doc.add_heading("Processing summary", level=1)
    summary = doc.add_table(rows=0, cols=2)
    summary.style = "Light Shading Accent 1"
    items = [
        ("Original registrations", str(original_count)),
        ("Duplicate registrations removed", str(duplicate_count)),
        ("Cleaned registrations saved", str(len(df))),
        ("Records requiring email review", str((df["Review_Status"] != "OK").sum())),
        ("Cleaned Excel output", OUTPUT_EXCEL.name),
    ]
    for label, value in items:
        cells = summary.add_row().cells
        cells[0].text = label
        cells[1].text = value

    doc.add_heading("Cleaning operations performed", level=1)
    operations = [
        "Removed extra spaces from column names and text fields.",
        "Converted email addresses to lowercase and marked missing email values for review.",
        "Standardized department labels (for example, Information Technology became IT).",
        "Removed duplicate registrations using email address, keeping the first occurrence.",
        "Sorted the final records by department and student name.",
        "Saved the cleaned data to a new Excel workbook so the original sample file remains unchanged.",
    ]
    for operation in operations:
        doc.add_paragraph(operation, style="List Bullet")

    doc.add_heading("Cleaned registration preview", level=1)
    preview = df.head(8)
    columns = ["Registration_ID", "Full_Name", "Email", "Department", "Year", "Review_Status"]
    columns = [col for col in columns if col in preview.columns]
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for i, col in enumerate(columns):
        table.rows[0].cells[i].text = col.replace("_", " ")
    for _, row in preview.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(columns):
            value = row[col]
            cells[i].text = "" if pd.isna(value) else str(value)

    doc.add_paragraph(
        "Note: The sample registration data is fictional and is provided only for demonstration."
    )
    doc.save(OUTPUT_WORD)

if __name__ == "__main__":
    cleaned, original_count, duplicate_count = clean_registration_data()
    create_word_report(cleaned, original_count, duplicate_count)
    print("Processing completed successfully!")
    print(f"Original registrations: {original_count}")
    print(f"Duplicate registrations removed: {duplicate_count}")
    print(f"Cleaned registrations saved: {len(cleaned)}")
    print(f"Excel output: {OUTPUT_EXCEL.name}")
    print(f"Word output: {OUTPUT_WORD.name}")
